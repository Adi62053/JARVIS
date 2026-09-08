import os
import re
import ctypes
from pathlib import Path


class FileSystemControl:
    """
    JARVIS V4 full-computer filesystem controller.

    Capabilities:
        - Search files
        - Search folders
        - Natural-language file search
        - Find files
        - Find folders
        - Inspect files
        - Inspect folders
        - Create files
        - Create folders
        - Read text files
        - Read PDF files
        - Read DOCX files
        - Read file by number
        - Open file by number
        - Close file by number
        - Open folder by number
        - Open folder by name
        - Close folder by number
        - Close folder by name
        - Read file by filename
        - Read file by direct path
        - Smart resume recognition

    Create behavior:
        - If no location is specified, creates on Desktop.
        - Multi-word names are preserved.
        - Existing files/folders are never silently overwritten.

    Safety:
        - Searches all available drives
        - Skips system/cache directories
        - Maximum 20 search results
        - Maximum 1 MB text/document read
        - Binary files rejected
        - Display output limited
        - File/folder closing uses normal Windows WM_CLOSE
    """

    MAX_RESULTS = 20
    MAX_READ_SIZE = 1_000_000
    MAX_DISPLAY_CHARS = 30_000

    # ==========================================
    # DIRECTORIES THAT SHOULD NOT BE CRAWLED
    # ==========================================

    SKIP_DIRECTORIES = {
        ".venv",
        "venv",
        "env",
        "__pycache__",
        ".git",
        ".svn",
        ".hg",
        "node_modules",
        "$recycle.bin",
        "recycle.bin",
        "system volume information",
        "windows",
        "program files",
        "program files (x86)",
        "programdata",
        "appdata",
    }

    # ==========================================
    # TEXT FILE EXTENSIONS
    # ==========================================

    TEXT_EXTENSIONS = {
        ".txt",
        ".py",
        ".js",
        ".jsx",
        ".ts",
        ".tsx",
        ".java",
        ".c",
        ".cpp",
        ".h",
        ".hpp",
        ".cs",
        ".go",
        ".rs",
        ".php",
        ".html",
        ".htm",
        ".css",
        ".scss",
        ".json",
        ".xml",
        ".yaml",
        ".yml",
        ".md",
        ".markdown",
        ".ini",
        ".cfg",
        ".conf",
        ".log",
        ".sql",
        ".bat",
        ".cmd",
        ".ps1",
        ".sh",
        ".env",
        ".gitignore",
        ".csv",
    }

    # ==========================================
    # DOCUMENT EXTENSIONS
    # ==========================================

    DOCUMENT_EXTENSIONS = {
        ".pdf",
        ".docx",
        ".doc",
    }

    # ==========================================
    # CONSTRUCTOR
    # ==========================================

    def __init__(self):

        self.last_search_results = []

        self.last_search_type = None

        self.last_read_file = None

        self.last_opened_file = None

        # Remember the last folder opened by JARVIS.
        self.last_opened_folder = None

    # ==========================================
    # AVAILABLE DRIVES
    # ==========================================

    def get_drives(self):

        drives = []

        for letter in "ABCDEFGHIJKLMNOPQRSTUVWXYZ":

            drive = Path(f"{letter}:\\")

            try:

                if drive.exists() and drive.is_dir():

                    drives.append(drive)

            except OSError:

                continue

        return drives

    # ==========================================
    # MAIN EXECUTOR
    # ==========================================

    def execute(self, command):

        if not command:

            return (
                "Sir, please tell me what filesystem "
                "operation you want."
            )

        command = str(command).lower().strip()

        command = self.normalize_command(command)

        # ======================================
        # CREATE FOLDER
        # ======================================

        create_folder = self.extract_create_folder(
            command
        )

        if create_folder is not None:

            if not create_folder:

                return (
                    "Sir, please tell me the name "
                    "of the folder you want me to create."
                )

            return self.create_folder(
                create_folder
            )

        # ======================================
        # CREATE FILE
        # ======================================

        create_file = self.extract_create_file(
            command
        )

        if create_file is not None:

            if not create_file:

                return (
                    "Sir, please tell me the name "
                    "of the file you want me to create."
                )

            return self.create_file(
                create_file
            )

        # ======================================
        # NATURAL LANGUAGE SEARCH
        # ======================================

        natural_search = self.extract_natural_search(
            command
        )

        if natural_search is not None:

            search_type, query = natural_search

            if search_type == "file":

                return self.search_files(query)

            if search_type == "folder":

                return self.search_folders(query)

        # ======================================
        # SEARCH FILES
        # ======================================

        if command in [
            "search file",
            "search files",
            "find file",
            "find files",
        ]:

            return self.search_files("")

        for prefix in [
            "search file ",
            "search files ",
            "find file ",
            "find files ",
        ]:

            if command.startswith(prefix):

                query = command[
                    len(prefix):
                ].strip()

                return self.search_files(query)

        # ======================================
        # SEARCH FOLDERS
        # ======================================

        if command in [
            "search folder",
            "search folders",
            "find folder",
            "find folders",
        ]:

            return self.search_folders("")

        for prefix in [
            "search folder ",
            "search folders ",
            "find folder ",
            "find folders ",
        ]:

            if command.startswith(prefix):

                query = command[
                    len(prefix):
                ].strip()

                return self.search_folders(query)

        # ======================================
        # INSPECT FILE
        # ======================================

        if command == "inspect file":

            return (
                "Sir, please specify the file "
                "you want me to inspect."
            )

        if command.startswith("inspect file "):

            file_reference = command[
                len("inspect file "):
            ].strip()

            return self.inspect_file(
                file_reference
            )

        for prefix in [
            "inspect my ",
            "inspect the ",
        ]:

            if command.startswith(prefix):

                file_reference = command[
                    len(prefix):
                ].strip()

                return self.inspect_file(
                    file_reference
                )

        # ======================================
        # INSPECT FOLDER
        # ======================================

        if command == "inspect folder":

            return (
                "Sir, please specify the folder "
                "you want me to inspect."
            )

        if command.startswith("inspect folder "):

            folder_reference = command[
                len("inspect folder "):
            ].strip()

            return self.inspect_folder(
                folder_reference
            )

        # ======================================
        # READ FILE
        # ======================================

        if command in [
            "read file",
            "read files",
        ]:

            if not self.last_search_results:

                return (
                    "Sir, there are no recent file "
                    "search results. Please search "
                    "for files first."
                )

            return (
                f"Sir, I found "
                f"{len(self.last_search_results)} "
                "files in the latest search. "
                "Please tell me the file number "
                "you want to read."
            )

        # ======================================
        # READ FILE NUMBER - INCOMPLETE
        # ======================================

        if command in [
            "read file number",
            "read files number",
            "read file no",
            "read files no",
        ]:

            return (
                "Sir, please specify the file number."
            )

        # ======================================
        # READ CURRENT FILE
        # ======================================

        current_file_commands = [

            "read this file",
            "read the file",
            "read this",
            "read the content",
            "read the content of this file",
            "read contents of this file",
            "read the contents of this file",
            "read content of this file",

        ]

        if command in current_file_commands:

            if self.last_read_file is None:

                return (
                    "Sir, there is no recently read "
                    "file. Please search for a file first."
                )

            if not self.last_read_file.exists():

                self.last_read_file = None

                return (
                    "Sir, the previously selected file "
                    "no longer exists."
                )

            return self.read_file(
                self.last_read_file
            )

        # ======================================
        # OPEN FILE NUMBER - INCOMPLETE
        # ======================================

        if command in [
            "open file number",
            "open files number",
            "open file no",
            "open files no",
        ]:

            return (
                "Sir, please specify the file number."
            )

        # ======================================
        # OPEN FILE BY NUMBER
        # ======================================

        for prefix in [
            "open file number ",
            "open files number ",
            "open file no ",
            "open files no ",
        ]:

            if command.startswith(prefix):

                number_text = command[
                    len(prefix):
                ].strip()

                if not number_text:

                    return (
                        "Sir, please specify "
                        "the file number."
                    )

                return self.open_file_by_number(
                    number_text
                )

        # ======================================
        # OPEN FOLDER NUMBER - INCOMPLETE
        # ======================================

        if command in [
            "open folder number",
            "open folders number",
            "open folder no",
            "open folders no",
        ]:

            return (
                "Sir, please specify the folder number."
            )

        # ======================================
        # OPEN FOLDER BY NUMBER
        # ======================================

        for prefix in [
            "open folder number ",
            "open folders number ",
            "open folder no ",
            "open folders no ",
        ]:

            if command.startswith(prefix):

                number_text = command[
                    len(prefix):
                ].strip()

                if not number_text:

                    return (
                        "Sir, please specify "
                        "the folder number."
                    )

                return self.open_folder_by_number(
                    number_text
                )

        # ======================================
        # OPEN FOLDER BY NAME
        # ======================================

        for prefix in [
            "open folder ",
            "open folders ",
        ]:

            if command.startswith(prefix):

                folder_reference = command[
                    len(prefix):
                ].strip()

                if not folder_reference:

                    return (
                        "Sir, please specify "
                        "the folder name."
                    )

                return self.open_folder_by_reference(
                    folder_reference
                )

        # ======================================
        # CLOSE FOLDER - INCOMPLETE
        # ======================================

        if command in [
            "close folder",
            "close folders",
            "close folder number",
            "close folders number",
            "close folder no",
            "close folders no",
        ]:

            return (
                "Sir, please specify the folder "
                "number or folder name to close."
            )

        # ======================================
        # CLOSE FOLDER BY NUMBER
        # ======================================

        for prefix in [
            "close folder number ",
            "close folders number ",
            "close folder no ",
            "close folders no ",
        ]:

            if command.startswith(prefix):

                number_text = command[
                    len(prefix):
                ].strip()

                if not number_text:

                    return (
                        "Sir, please specify "
                        "the folder number."
                    )

                return self.close_folder_by_number(
                    number_text
                )

        # ======================================
        # CLOSE FOLDER BY NAME
        # ======================================

        for prefix in [
            "close folder ",
            "close folders ",
        ]:

            if command.startswith(prefix):

                folder_reference = command[
                    len(prefix):
                ].strip()

                if not folder_reference:

                    return (
                        "Sir, please specify "
                        "the folder name."
                    )

                return self.close_folder_by_reference(
                    folder_reference
                )

        # ======================================
        # CLOSE FILE - INCOMPLETE
        # ======================================

        if command in [
            "close file",
            "close files",
            "close file number",
            "close files number",
            "close file no",
            "close files no",
        ]:

            return (
                "Sir, please specify the file "
                "number to close."
            )

        # ======================================
        # CLOSE FILE BY NUMBER
        # ======================================

        for prefix in [
            "close file number ",
            "close files number ",
            "close file no ",
            "close files no ",
            "close file ",
            "close files ",
        ]:

            if command.startswith(prefix):

                number_text = command[
                    len(prefix):
                ].strip()

                if not number_text:

                    return (
                        "Sir, please specify the file "
                        "number to close."
                    )

                return self.close_file_by_number(
                    number_text
                )

        # ======================================
        # READ FILE NUMBER
        # ======================================

        for prefix in [
            "read file number ",
            "read files number ",
            "read file no ",
            "read files no ",
        ]:

            if command.startswith(prefix):

                number_text = command[
                    len(prefix):
                ].strip()

                if not number_text:

                    return (
                        "Sir, please specify "
                        "the file number."
                    )

                return self.read_file_by_number(
                    number_text
                )

        # ======================================
        # FILE NUMBER
        # ======================================

        if command in [
            "file number",
            "file no",
        ]:

            return (
                "Sir, please specify the file number."
            )

        for prefix in [
            "file number ",
            "file no ",
        ]:

            if command.startswith(prefix):

                number_text = command[
                    len(prefix):
                ].strip()

                if not number_text:

                    return (
                        "Sir, please specify "
                        "the file number."
                    )

                return self.read_file_by_number(
                    number_text
                )

        # ======================================
        # NATURAL LANGUAGE READ
        # ======================================

        natural_read = self.extract_natural_read(
            command
        )

        if natural_read:

            return self.read_file_by_reference(
                natural_read
            )

        # ======================================
        # READ FILE BY NAME
        # ======================================

        for prefix in [
            "read file ",
            "read files ",
        ]:

            if command.startswith(prefix):

                file_reference = command[
                    len(prefix):
                ].strip()

                return self.read_file_by_reference(
                    file_reference
                )

        return (
            "I don't have a filesystem tool "
            "for that command yet."
        )

    # ==========================================
    # CREATE COMMAND EXTRACTION
    # ==========================================

    def extract_create_folder(self, command):

        prefixes = [
            "create folder ",
            "create folders ",
            "make folder ",
            "make folders ",
            "new folder ",
            "new folders ",
        ]

        if command in [
            "create folder",
            "create folders",
            "make folder",
            "make folders",
            "new folder",
            "new folders",
        ]:

            return ""

        for prefix in prefixes:

            if command.startswith(prefix):

                name = command[
                    len(prefix):
                ].strip()

                return self.clean_create_target(
                    name
                )

        return None

    def extract_create_file(self, command):

        prefixes = [
            "create file ",
            "create files ",
            "make file ",
            "make files ",
            "new file ",
            "new files ",
        ]

        if command in [
            "create file",
            "create files",
            "make file",
            "make files",
            "new file",
            "new files",
        ]:

            return ""

        for prefix in prefixes:

            if command.startswith(prefix):

                name = command[
                    len(prefix):
                ].strip()

                return self.clean_create_target(
                    name
                )

        return None

    # ==========================================
    # CLEAN CREATE TARGET
    # ==========================================

    def clean_create_target(self, target):

        target = str(
            target
        ).strip()

        target = target.strip(
            "\"'"
        )

        target = re.sub(
            r"^(called|named|name)\s+",
            "",
            target,
            flags=re.IGNORECASE,
        )

        return target.strip()

    # ==========================================
    # DEFAULT CREATE LOCATION
    # ==========================================

    def get_default_create_directory(self):

        desktop = (
            Path.home()
            / "Desktop"
        )

        if desktop.exists() and desktop.is_dir():

            return desktop

        return Path.cwd()

    # ==========================================
    # SPECIAL WINDOWS LOCATIONS
    # ==========================================

    def resolve_create_target(self, target):

        target = str(
            target
        ).strip()

        location_match = re.search(
            r"\s+(?:on|in|inside)\s+(desktop|documents|downloads)$",
            target,
            flags=re.IGNORECASE,
        )

        if location_match:

            location_name = (
                location_match.group(1).lower()
            )

            name = target[
                :location_match.start()
            ].strip()

            if location_name == "desktop":

                parent = (
                    Path.home()
                    / "Desktop"
                )

            elif location_name == "documents":

                parent = (
                    Path.home()
                    / "Documents"
                )

            else:

                parent = (
                    Path.home()
                    / "Downloads"
                )

            return parent, name

        if (
            ":" in target[:3]
            or "\\" in target
            or "/" in target
        ):

            explicit_path = Path(
                target
            )

            if explicit_path.is_absolute():

                return (
                    explicit_path.parent,
                    explicit_path.name,
                )

        return (
            self.get_default_create_directory(),
            target,
        )

    # ==========================================
    # VALIDATE CREATE NAME
    # ==========================================

    def validate_create_name(self, name):

        if not name:

            return (
                False,
                "Sir, the name cannot be empty."
            )

        if name in [
            ".",
            "..",
        ]:

            return (
                False,
                "Sir, that is not a valid name."
            )

        invalid_characters = (
            '<>:"/\\|?*'
        )

        if any(
            character in name
            for character in invalid_characters
        ):

            return (
                False,
                "Sir, that name contains characters "
                "that are not valid in a Windows filename."
            )

        if name.endswith(
            " "
        ) or name.endswith(
            "."
        ):

            return (
                False,
                "Sir, a Windows file or folder name "
                "cannot end with a space or period."
            )

        return (
            True,
            None
        )

    # ==========================================
    # CREATE FOLDER
    # ==========================================

    def create_folder(self, target):

        parent, name = (
            self.resolve_create_target(
                target
            )
        )

        valid, error = (
            self.validate_create_name(
                name
            )
        )

        if not valid:

            return error

        try:

            parent = parent.expanduser()

            if not parent.exists():

                return (
                    f"Sir, the parent folder "
                    f"'{parent}' does not exist."
                )

            if not parent.is_dir():

                return (
                    f"Sir, '{parent}' is not a folder."
                )

            new_folder = (
                parent / name
            )

            if new_folder.exists():

                return (
                    f"Sir, a file or folder named "
                    f"'{name}' already exists at "
                    f"'{parent}'."
                )

            new_folder.mkdir()

            return (
                f"Sir, I created the folder "
                f"'{name}'."
                f"|||VOICE|||"
                f"Sir, I created the folder {name}."
            )

        except PermissionError:

            return (
                "Sir, I don't have permission "
                "to create that folder."
            )

        except OSError as e:

            print(
                "JARVIS folder creation error:",
                e
            )

            return (
                f"Sir, I couldn't create the "
                f"folder '{name}'."
            )

    # ==========================================
    # CREATE FILE
    # ==========================================

    def create_file(self, target):

        parent, name = (
            self.resolve_create_target(
                target
            )
        )

        valid, error = (
            self.validate_create_name(
                name
            )
        )

        if not valid:

            return error

        try:

            parent = parent.expanduser()

            if not parent.exists():

                return (
                    f"Sir, the parent folder "
                    f"'{parent}' does not exist."
                )

            if not parent.is_dir():

                return (
                    f"Sir, '{parent}' is not a folder."
                )

            new_file = (
                parent / name
            )

            if new_file.exists():

                return (
                    f"Sir, a file or folder named "
                    f"'{name}' already exists at "
                    f"'{parent}'."
                )

            with open(
                new_file,
                "x",
                encoding="utf-8",
            ):
                pass

            return (
                f"Sir, I created the file "
                f"'{name}'."
                f"|||VOICE|||"
                f"Sir, I created the file {name}."
            )

        except PermissionError:

            return (
                "Sir, I don't have permission "
                "to create that file."
            )

        except OSError as e:

            print(
                "JARVIS file creation error:",
                e
            )

            return (
                f"Sir, I couldn't create "
                f"the file '{name}'."
            )

    # ==========================================
    # COMMAND NORMALIZATION
    # ==========================================

    def normalize_command(self, command):

        command = str(
            command
        ).lower().strip()

        replacements = {

            "open file no. ": "open file number ",
            "open files no. ": "open file number ",

            "read file no. ": "read file number ",
            "read files no. ": "read file number ",

            "file no. ": "file number ",

            "close file no. ": "close file number ",
            "close files no. ": "close files number ",

            "open folder no. ": "open folder number ",
            "open folders no. ": "open folders number ",

            "close folder no. ": "close folder number ",
            "close folders no. ": "close folders number ",

        }

        for spoken, normal in replacements.items():

            if command.startswith(spoken):

                command = (
                    normal
                    + command[len(spoken):]
                )

                break

        return command.strip()

    # ==========================================
    # NATURAL LANGUAGE SEARCH COMMANDS
    # ==========================================

    def extract_natural_search(self, command):

        resume_phrases = [

            "search resume",
            "search my resume",
            "search my resume file",
            "search resume file",

            "find resume",
            "find my resume",
            "find my resume file",
            "find resume file",

        ]

        if command in resume_phrases:

            return "file", "resume"

        if command in [

            "search res",
            "search resu",
            "search resum",

            "find res",
            "find resu",
            "find resum",

            "search my res",
            "search my resu",
            "search my resum",

            "find my res",
            "find my resu",
            "find my resum",

        ]:

            return "file", "resume"

        pdf_phrases = [

            "search pdf",
            "search pdf file",
            "search pdf files",

            "find pdf",
            "find pdf file",
            "find pdf files",

            "search for pdf",
            "search for pdf files",

            "find my pdf",
            "find my pdf files",

        ]

        if command in pdf_phrases:

            return "file", ".pdf"

        docx_phrases = [

            "search docx",
            "search docx file",
            "search docx files",

            "find docx",
            "find docx file",
            "find docx files",

            "search word document",
            "search word documents",

            "find word document",
            "find word documents",

        ]

        if command in docx_phrases:

            return "file", ".docx"

        document_phrases = [

            "search document",
            "search documents",

            "find document",
            "find documents",

            "search my documents",
            "find my documents",

        ]

        if command in document_phrases:

            return "file", ""

        prefixes = [

            "search my ",
            "search for my ",
            "search for ",
            "search ",

            "find my ",
            "find ",

        ]

        for prefix in prefixes:

            if command.startswith(prefix):

                query = command[
                    len(prefix):
                ].strip()

                if not query:

                    return None

                blocked = [
                    "file",
                    "files",
                    "folder",
                    "folders",
                ]

                if query in blocked:

                    return None

                resume_words = [
                    "resume",
                    "resum",
                    "resu",
                    "res",
                ]

                if query in resume_words:

                    return "file", "resume"

                if (
                    "resume" in query
                    or query.startswith("resum")
                    or query == "res"
                ):

                    return "file", "resume"

                if "pdf" in query:

                    return "file", ".pdf"

                if (
                    "docx" in query
                    or "word document" in query
                    or "word documents" in query
                ):

                    return "file", ".docx"

                if (
                    query.endswith("folder")
                    or query.endswith("folders")
                ):

                    query = re.sub(
                        r"\bfolders?\b",
                        "",
                        query
                    ).strip()

                    return "folder", query

                return "file", query

        return None

    # ==========================================
    # NATURAL LANGUAGE READ
    # ==========================================

    def extract_natural_read(self, command):

        prefixes = [

            "read my ",
            "read the ",
            "read ",

        ]

        for prefix in prefixes:

            if command.startswith(prefix):

                reference = command[
                    len(prefix):
                ].strip()

                if not reference:

                    return None

                if reference.startswith(
                    "file number "
                ):

                    return None

                if reference.startswith(
                    "file no "
                ):

                    return None

                if reference in [

                    "resume",
                    "resume file",
                    "my resume",
                    "res",
                    "resu",
                    "resum",

                ]:

                    return "resume"

                return reference

        return None

    # ==========================================
    # SEARCH FILES
    # ==========================================

    def search_files(self, query):

        query = self.normalize_search_query(
            query
        )

        self.last_search_results = []

        self.last_search_type = "file"

        results = []

        for drive in self.get_drives():

            try:

                for root, dirs, files in os.walk(
                    drive,
                    topdown=True,
                ):

                    dirs[:] = [
                        directory
                        for directory in dirs
                        if directory.lower()
                        not in self.SKIP_DIRECTORIES
                    ]

                    for filename in files:

                        filename_lower = (
                            filename.lower()
                        )

                        if (
                            query
                            and query not in filename_lower
                        ):

                            continue

                        full_path = (
                            Path(root) / filename
                        )

                        try:

                            if not full_path.is_file():

                                continue

                        except OSError:

                            continue

                        results.append(
                            full_path
                        )

            except (
                PermissionError,
                OSError,
            ):

                continue

        if query and results:

            query_lower = query.lower()

            def result_score(path):

                filename = path.name.lower()

                if filename == query_lower:

                    return 0

                if filename.startswith(
                    query_lower
                ):

                    return 1

                return 2

            results.sort(
                key=result_score
            )

        results = results[
            :self.MAX_RESULTS
        ]

        self.last_search_results = list(
            results
        )

        if not results:

            if query:

                return (
                    f"Sir, I couldn't find any files "
                    f"matching '{query}' on the "
                    "available drives."
                )

            return (
                "Sir, I couldn't find any files "
                "on the available drives."
            )

        console_lines = []

        console_lines.append(
            f"Found {len(results)} file(s) "
            "on the available drives:"
        )

        console_lines.append("")

        for index, path in enumerate(
            results,
            start=1
        ):

            console_lines.append(
                f"{index}. {path}"
            )

        console_lines.append("")

        console_lines.append(
            "Say 'read file number N' "
            "to read a file."
        )

        console_lines.append(
            "Say 'open file number N' "
            "to open a file."
        )

        console_lines.append(
            "Say 'close file number N' "
            "to close an opened file."
        )

        console_output = "\n".join(
            console_lines
        )

        voice_output = (
            f"Sir, I found {len(results)} files. "
            "Please tell me which file number "
            "you want to read or open."
        )

        return (
            f"{console_output}"
            f"|||VOICE|||"
            f"{voice_output}"
        )

    # ==========================================
    # SEARCH FOLDERS
    # ==========================================

    def search_folders(self, query):

        query = self.normalize_search_query(
            query
        )

        self.last_search_results = []

        self.last_search_type = "folder"

        results = []

        for drive in self.get_drives():

            try:

                for root, dirs, files in os.walk(
                    drive,
                    topdown=True,
                ):

                    dirs[:] = [
                        directory
                        for directory in dirs
                        if directory.lower()
                        not in self.SKIP_DIRECTORIES
                    ]

                    for dirname in list(dirs):

                        if (
                            query
                            and query not in dirname.lower()
                        ):

                            continue

                        full_path = (
                            Path(root) / dirname
                        )

                        try:

                            if not full_path.is_dir():

                                continue

                        except OSError:

                            continue

                        results.append(
                            full_path
                        )

            except (
                PermissionError,
                OSError,
            ):

                continue

        if query and results:

            query_lower = query.lower()

            def folder_score(path):

                name = path.name.lower()

                if name == query_lower:

                    return 0

                if name.startswith(
                    query_lower
                ):

                    return 1

                return 2

            results.sort(
                key=folder_score
            )

        results = results[
            :self.MAX_RESULTS
        ]

        self.last_search_results = list(
            results
        )

        if not results:

            if query:

                return (
                    f"Sir, I couldn't find any "
                    f"folders matching '{query}' "
                    "on the available drives."
                )

            return (
                "Sir, I couldn't find any folders "
                "on the available drives."
            )

        console_lines = []

        console_lines.append(
            f"Found {len(results)} folder(s) "
            "on the available drives:"
        )

        console_lines.append("")

        for index, path in enumerate(
            results,
            start=1
        ):

            console_lines.append(
                f"{index}. {path}"
            )

        console_lines.append("")

        console_lines.append(
            "Say 'open folder number N' "
            "to open a folder."
        )

        console_lines.append(
            "Say 'close folder number N' "
            "to close a folder."
        )

        console_output = "\n".join(
            console_lines
        )

        voice_output = (
            f"Sir, I found {len(results)} folders. "
            "Please tell me the folder number "
            "you want to open or close."
        )

        return (
            f"{console_output}"
            f"|||VOICE|||"
            f"{voice_output}"
        )

    # ==========================================
    # READ FILE BY NUMBER
    # ==========================================

    def read_file_by_number(
        self,
        number_text
    ):

        number = self.parse_number(
            number_text
        )

        if number is None:

            return (
                "Sir, I couldn't understand "
                "the file number."
            )

        if not self.last_search_results:

            return (
                "Sir, there are no recent file "
                "search results. Please search "
                "for files first."
            )

        if self.last_search_type != "file":

            return (
                "Sir, the latest search contains "
                "folders, not files. Please search "
                "for a file first."
            )

        if (
            number < 1
            or number > len(
                self.last_search_results
            )
        ):

            return (
                f"Sir, file number {number} "
                "does not exist in the latest "
                "search results."
            )

        path = self.last_search_results[
            number - 1
        ]

        if not path.exists():

            self.last_search_results = [
                item
                for item in self.last_search_results
                if item.exists()
            ]

            return (
                "Sir, that file was removed "
                "after the search."
            )

        if not path.is_file():

            return (
                f"Sir, item number {number} "
                "is not a file."
            )

        return self.read_file(path)

    # ==========================================
    # OPEN FILE BY NUMBER
    # ==========================================

    def open_file_by_number(
        self,
        number_text
    ):

        number = self.parse_number(
            number_text
        )

        if number is None:

            return (
                "Sir, I couldn't understand "
                "the file number."
            )

        if not self.last_search_results:

            return (
                "Sir, there are no recent search "
                "results. Please search for a file first."
            )

        if self.last_search_type != "file":

            return (
                "Sir, the latest search contains "
                "folders, not files."
            )

        if (
            number < 1
            or number > len(
                self.last_search_results
            )
        ):

            return (
                f"Sir, file number {number} "
                "does not exist in the latest "
                "search results."
            )

        path = self.last_search_results[
            number - 1
        ]

        if not path.exists():

            return (
                "Sir, that file no longer exists."
            )

        if not path.is_file():

            return (
                "Sir, that item is not a file."
            )

        try:

            os.startfile(
                str(path)
            )

        except PermissionError:

            return (
                "Sir, I don't have permission "
                "to open that file."
            )

        except OSError as e:

            print(
                "JARVIS file open error:",
                e
            )

            return (
                f"Sir, I couldn't open "
                f"'{path.name}'."
            )

        self.last_read_file = path
        self.last_opened_file = path

        return (
            f"Sir, I opened "
            f"'{path.name}'."
            "|||VOICE|||"
            f"Sir, I opened {path.name}."
        )

    # ==========================================
    # CLOSE FILE BY NUMBER
    # ==========================================

    def close_file_by_number(
        self,
        number_text
    ):

        number = self.parse_number(
            number_text
        )

        if number is None:

            return (
                "Sir, I couldn't understand "
                "the file number to close."
            )

        if not self.last_search_results:

            return (
                "Sir, there are no recent search "
                "results. Please search for the "
                "file first."
            )

        if self.last_search_type != "file":

            return (
                "Sir, the latest search contains "
                "folders, not files."
            )

        if (
            number < 1
            or number > len(
                self.last_search_results
            )
        ):

            return (
                f"Sir, file number {number} "
                "does not exist in the latest "
                "search results."
            )

        path = self.last_search_results[
            number - 1
        ]

        if not path.exists():

            return (
                "Sir, that file no longer exists."
            )

        if not path.is_file():

            return (
                "Sir, that item is not a file."
            )

        closed = self.close_windows_for_file(
            path
        )

        if closed:

            if self.last_opened_file == path:

                self.last_opened_file = None

            return (
                f"Sir, I closed "
                f"'{path.name}'."
                "|||VOICE|||"
                f"Sir, I closed {path.name}."
            )

        return (
            f"Sir, I couldn't find an open "
            f"window for '{path.name}'."
            "|||VOICE|||"
            f"Sir, I couldn't find an open "
            f"window for {path.name}."
        )

    # ==========================================
    # CLOSE WINDOWS ASSOCIATED WITH FILE
    # ==========================================

    def close_windows_for_file(
        self,
        path
    ):

        if os.name != "nt":

            return False

        try:

            user32 = ctypes.windll.user32

            WM_CLOSE = 0x0010

            closed = False

            target_name = path.name.lower()

            target_stem = path.stem.lower()

            EnumWindowsProc = ctypes.WINFUNCTYPE(
                ctypes.c_bool,
                ctypes.c_void_p,
                ctypes.c_void_p,
            )

            def enum_windows_proc(
                hwnd,
                lparam
            ):

                nonlocal closed

                try:

                    if not user32.IsWindowVisible(
                        hwnd
                    ):

                        return True

                    title_length = (
                        user32.GetWindowTextLengthW(
                            hwnd
                        )
                    )

                    if title_length <= 0:

                        return True

                    buffer = (
                        ctypes.create_unicode_buffer(
                            title_length + 1
                        )
                    )

                    user32.GetWindowTextW(
                        hwnd,
                        buffer,
                        title_length + 1,
                    )

                    title = (
                        buffer.value
                        .strip()
                        .lower()
                    )

                    if not title:

                        return True

                    matches_file = (
                        target_name in title
                    )

                    matches_stem = (
                        target_stem
                        and target_stem in title
                    )

                    if (
                        matches_file
                        or matches_stem
                    ):

                        user32.PostMessageW(
                            hwnd,
                            WM_CLOSE,
                            0,
                            0,
                        )

                        closed = True

                except Exception as e:

                    print(
                        "JARVIS window close "
                        "inspection error:",
                        e
                    )

                return True

            callback = EnumWindowsProc(
                enum_windows_proc
            )

            user32.EnumWindows(
                callback,
                0,
            )

            return closed

        except Exception as e:

            print(
                "JARVIS close window error:",
                e
            )

            return False

    # ==========================================
    # OPEN FOLDER BY NUMBER
    # ==========================================

    def open_folder_by_number(
        self,
        number_text
    ):

        number = self.parse_number(
            number_text
        )

        if number is None:

            return (
                "Sir, I couldn't understand "
                "the folder number."
            )

        if not self.last_search_results:

            return (
                "Sir, there are no recent search "
                "results. Please search for a folder first."
            )

        if self.last_search_type != "folder":

            return (
                "Sir, the latest search contains "
                "files, not folders."
            )

        if (
            number < 1
            or number > len(
                self.last_search_results
            )
        ):

            return (
                f"Sir, folder number {number} "
                "does not exist in the latest "
                "search results."
            )

        path = self.last_search_results[
            number - 1
        ]

        return self.open_folder_path(
            path
        )

    # ==========================================
    # OPEN FOLDER BY NAME / PATH
    # ==========================================

    def open_folder_by_reference(
        self,
        folder_reference
    ):

        folder_reference = (
            self.normalize_file_reference(
                folder_reference
            )
        )

        if not folder_reference:

            return (
                "Sir, please specify "
                "the folder name."
            )

        possible_path = Path(
            folder_reference
        )

        try:

            if (
                possible_path.exists()
                and possible_path.is_dir()
            ):

                return self.open_folder_path(
                    possible_path
                )

        except OSError:

            pass

        for path in self.last_search_results:

            try:

                if (
                    path.exists()
                    and path.is_dir()
                    and path.name.lower()
                    == folder_reference.lower()
                ):

                    return self.open_folder_path(
                        path
                    )

            except OSError:

                continue

        matches = (
            self.search_matching_folders(
                folder_reference
            )
        )

        if len(matches) == 1:

            return self.open_folder_path(
                matches[0]
            )

        if len(matches) > 1:

            self.last_search_results = (
                matches[
                    :self.MAX_RESULTS
                ]
            )

            self.last_search_type = "folder"

            console_lines = [
                f"Sir, I found {len(matches)} "
                f"folders matching "
                f"'{folder_reference}':",
                "",
            ]

            for index, path in enumerate(
                self.last_search_results,
                start=1
            ):

                console_lines.append(
                    f"{index}. {path}"
                )

            console_lines.append("")
            console_lines.append(
                "Say 'open folder number N' "
                "to open one."
            )

            return (
                "\n".join(console_lines)
                + "|||VOICE|||"
                + f"Sir, I found {len(matches)} "
                "matching folders. Please tell "
                "me the folder number."
            )

        return (
            f"Sir, I couldn't find a folder "
            f"named '{folder_reference}'."
        )

    # ==========================================
    # OPEN FOLDER PATH
    # ==========================================

    def open_folder_path(
        self,
        path
    ):

        path = Path(
            path
        )

        if not path.exists():

            return (
                f"Sir, the folder '{path}' "
                "no longer exists."
            )

        if not path.is_dir():

            return (
                f"Sir, '{path}' is not a folder."
            )

        try:

            os.startfile(
                str(path)
            )

        except PermissionError:

            return (
                "Sir, I don't have permission "
                "to open that folder."
            )

        except OSError as e:

            print(
                "JARVIS folder open error:",
                e
            )

            return (
                f"Sir, I couldn't open "
                f"'{path.name}'."
            )

        self.last_opened_folder = path

        return (
            f"Sir, I opened the folder "
            f"'{path.name}'."
            "|||VOICE|||"
            f"Sir, I opened the folder {path.name}."
        )

    # ==========================================
    # CLOSE FOLDER BY NUMBER
    # ==========================================

    def close_folder_by_number(
        self,
        number_text
    ):

        number = self.parse_number(
            number_text
        )

        if number is None:

            return (
                "Sir, I couldn't understand "
                "the folder number to close."
            )

        if not self.last_search_results:

            return (
                "Sir, there are no recent search "
                "results. Please search for a "
                "folder first."
            )

        if self.last_search_type != "folder":

            return (
                "Sir, the latest search contains "
                "files, not folders."
            )

        if (
            number < 1
            or number > len(
                self.last_search_results
            )
        ):

            return (
                f"Sir, folder number {number} "
                "does not exist in the latest "
                "search results."
            )

        path = self.last_search_results[
            number - 1
        ]

        if not path.exists():

            return (
                "Sir, that folder no longer exists."
            )

        if not path.is_dir():

            return (
                "Sir, that item is not a folder."
            )

        closed = self.close_windows_for_folder(
            path
        )

        if closed:

            if self.last_opened_folder == path:

                self.last_opened_folder = None

            return (
                f"Sir, I closed the folder "
                f"'{path.name}'."
                "|||VOICE|||"
                f"Sir, I closed the folder "
                f"{path.name}."
            )

        return (
            f"Sir, I couldn't find an open "
            f"window for the folder "
            f"'{path.name}'."
            "|||VOICE|||"
            f"Sir, I couldn't find an open "
            f"window for the folder "
            f"{path.name}."
        )

    # ==========================================
    # CLOSE FOLDER BY NAME / PATH
    # ==========================================

    def close_folder_by_reference(
        self,
        folder_reference
    ):

        folder_reference = (
            self.normalize_file_reference(
                folder_reference
            )
        )

        if not folder_reference:

            return (
                "Sir, please specify "
                "the folder name."
            )

        # --------------------------------------
        # DIRECT PATH
        # --------------------------------------

        possible_path = Path(
            folder_reference
        )

        try:

            if (
                possible_path.exists()
                and possible_path.is_dir()
            ):

                return self.close_folder_path(
                    possible_path
                )

        except OSError:

            pass

        # --------------------------------------
        # LAST OPENED FOLDER
        # --------------------------------------

        if self.last_opened_folder is not None:

            try:

                if (
                    self.last_opened_folder.exists()
                    and self.last_opened_folder.is_dir()
                    and self.last_opened_folder.name.lower()
                    == folder_reference.lower()
                ):

                    return self.close_folder_path(
                        self.last_opened_folder
                    )

            except OSError:

                pass

        # --------------------------------------
        # CURRENT SEARCH RESULTS
        # --------------------------------------

        exact_matches = []

        for path in self.last_search_results:

            try:

                if (
                    path.exists()
                    and path.is_dir()
                    and path.name.lower()
                    == folder_reference.lower()
                ):

                    exact_matches.append(
                        path
                    )

            except OSError:

                continue

        if len(exact_matches) == 1:

            return self.close_folder_path(
                exact_matches[0]
            )

        if len(exact_matches) > 1:

            self.last_search_results = (
                exact_matches[
                    :self.MAX_RESULTS
                ]
            )

            self.last_search_type = "folder"

            return self.format_folder_close_results(
                self.last_search_results,
                folder_reference
            )

        # --------------------------------------
        # SEARCH FOLDER BY NAME
        # --------------------------------------

        matches = (
            self.search_matching_folders(
                folder_reference
            )
        )

        if len(matches) == 1:

            return self.close_folder_path(
                matches[0]
            )

        if len(matches) > 1:

            self.last_search_results = (
                matches[
                    :self.MAX_RESULTS
                ]
            )

            self.last_search_type = "folder"

            return self.format_folder_close_results(
                self.last_search_results,
                folder_reference
            )

        return (
            f"Sir, I couldn't find a folder "
            f"named '{folder_reference}'."
        )

    # ==========================================
    # CLOSE FOLDER PATH
    # ==========================================

    def close_folder_path(
        self,
        path
    ):

        path = Path(
            path
        )

        if not path.exists():

            return (
                f"Sir, the folder '{path}' "
                "no longer exists."
            )

        if not path.is_dir():

            return (
                f"Sir, '{path}' is not a folder."
            )

        closed = self.close_windows_for_folder(
            path
        )

        if closed:

            if self.last_opened_folder == path:

                self.last_opened_folder = None

            return (
                f"Sir, I closed the folder "
                f"'{path.name}'."
                "|||VOICE|||"
                f"Sir, I closed the folder "
                f"{path.name}."
            )

        return (
            f"Sir, I couldn't find an open "
            f"window for the folder "
            f"'{path.name}'."
            "|||VOICE|||"
            f"Sir, I couldn't find an open "
            f"window for the folder "
            f"{path.name}."
        )

    # ==========================================
    # FORMAT FOLDER CLOSE RESULTS
    # ==========================================

    def format_folder_close_results(
        self,
        results,
        query=""
    ):

        console_lines = []

        console_lines.append(
            f"Sir, I found {len(results)} "
            f"folders matching '{query}':"
        )

        console_lines.append("")

        for index, path in enumerate(
            results,
            start=1
        ):

            console_lines.append(
                f"{index}. {path}"
            )

        console_lines.append("")

        console_lines.append(
            "Say 'close folder number N' "
            "to close one."
        )

        console_output = "\n".join(
            console_lines
        )

        voice_output = (
            f"Sir, I found {len(results)} "
            "matching folders. Please tell "
            "me the folder number to close."
        )

        return (
            f"{console_output}"
            f"|||VOICE|||"
            f"{voice_output}"
        )

    # ==========================================
    # CLOSE WINDOWS ASSOCIATED WITH FOLDER
    # ==========================================

    def close_windows_for_folder(
        self,
        path
    ):

        if os.name != "nt":

            return False

        path = Path(
            path
        )

        try:

            user32 = ctypes.windll.user32

            WM_CLOSE = 0x0010

            closed = False

            target_name = (
                path.name
                .strip()
                .lower()
            )

            # Windows Explorer window classes.
            explorer_classes = {
                "cabinetwclass",
                "explorewclass",
            }

            EnumWindowsProc = ctypes.WINFUNCTYPE(
                ctypes.c_bool,
                ctypes.c_void_p,
                ctypes.c_void_p,
            )

            def enum_windows_proc(
                hwnd,
                lparam
            ):

                nonlocal closed

                try:

                    # ----------------------------------
                    # Only visible windows.
                    # ----------------------------------

                    if not user32.IsWindowVisible(
                        hwnd
                    ):

                        return True

                    # ----------------------------------
                    # Identify the actual Windows
                    # Explorer window class.
                    # ----------------------------------

                    class_buffer = (
                        ctypes.create_unicode_buffer(
                            256
                        )
                    )

                    user32.GetClassNameW(
                        hwnd,
                        class_buffer,
                        256,
                    )

                    window_class = (
                        class_buffer.value
                        .strip()
                        .lower()
                    )

                    if (
                        window_class
                        not in explorer_classes
                    ):

                        return True

                    # ----------------------------------
                    # Read Explorer window title.
                    # ----------------------------------

                    title_length = (
                        user32.GetWindowTextLengthW(
                            hwnd
                        )
                    )

                    if title_length <= 0:

                        return True

                    title_buffer = (
                        ctypes.create_unicode_buffer(
                            title_length + 1
                        )
                    )

                    user32.GetWindowTextW(
                        hwnd,
                        title_buffer,
                        title_length + 1,
                    )

                    title = (
                        title_buffer.value
                        .strip()
                        .lower()
                    )

                    if not title:

                        return True

                    # ----------------------------------
                    # Windows Explorer normally uses
                    # the folder name as the title.
                    #
                    # We support:
                    #
                    # test jar
                    # test jar - file explorer
                    # file explorer - test jar
                    # ----------------------------------

                    exact_match = (
                        title == target_name
                    )

                    suffix_match = (
                        title.endswith(
                            f" - {target_name}"
                        )
                    )

                    prefix_match = (
                        title.startswith(
                            f"{target_name} - "
                        )
                    )

                    # ----------------------------------
                    # Also handle a title containing
                    # the folder name when Explorer
                    # adds additional text.
                    # ----------------------------------

                    contains_match = (
                        target_name
                        and target_name in title
                    )

                    if (
                        exact_match
                        or suffix_match
                        or prefix_match
                        or contains_match
                    ):

                        user32.PostMessageW(
                            hwnd,
                            WM_CLOSE,
                            0,
                            0,
                        )

                        closed = True

                except Exception as e:

                    print(
                        "JARVIS folder window "
                        "inspection error:",
                        e
                    )

                return True

            callback = EnumWindowsProc(
                enum_windows_proc
            )

            user32.EnumWindows(
                callback,
                0,
            )

            return closed

        except Exception as e:

            print(
                "JARVIS folder close error:",
                e
            )

            return False

    # ==========================================
    # SEARCH MATCHING FOLDERS
    # ==========================================

    def search_matching_folders(
        self,
        query
    ):

        query = self.normalize_search_query(
            query
        )

        results = []

        for drive in self.get_drives():

            try:

                for root, dirs, files in os.walk(
                    drive,
                    topdown=True,
                ):

                    dirs[:] = [
                        directory
                        for directory in dirs
                        if directory.lower()
                        not in self.SKIP_DIRECTORIES
                    ]

                    for dirname in dirs:

                        if (
                            query
                            not in dirname.lower()
                        ):

                            continue

                        path = (
                            Path(root)
                            / dirname
                        )

                        try:

                            if path.is_dir():

                                results.append(
                                    path
                                )

                        except OSError:

                            continue

            except (
                PermissionError,
                OSError,
            ):

                continue

        if results:

            query_lower = query.lower()

            results.sort(
                key=lambda p: (
                    0
                    if p.name.lower()
                    == query_lower
                    else 1
                    if p.name.lower()
                    .startswith(query_lower)
                    else 2
                )
            )

        return results[
            :self.MAX_RESULTS
        ]

    # ==========================================
    # READ FILE BY REFERENCE
    # ==========================================

    def read_file_by_reference(
        self,
        file_reference
    ):

        file_reference = (
            self.normalize_file_reference(
                file_reference
            )
        )

        if not file_reference:

            return (
                "Sir, please tell me "
                "which file you want me to read."
            )

        if self.is_resume_reference(
            file_reference
        ):

            matches = (
                self.search_matching_files(
                    "resume"
                )
            )

            if len(matches) == 1:

                return self.read_file(
                    matches[0]
                )

            if len(matches) > 1:

                self.last_search_results = (
                    matches[
                        :self.MAX_RESULTS
                    ]
                )

                self.last_search_type = "file"

                return self.format_search_results(
                    self.last_search_results,
                    "resume"
                )

        for path in self.last_search_results:

            try:

                if (
                    path.exists()
                    and path.is_file()
                    and path.name.lower()
                    == file_reference.lower()
                ):

                    return self.read_file(
                        path
                    )

            except OSError:

                continue

        possible_path = Path(
            file_reference
        )

        try:

            if possible_path.exists():

                if possible_path.is_file():

                    return self.read_file(
                        possible_path
                    )

                return (
                    "Sir, that path is a folder, "
                    "not a file."
                )

        except OSError:

            pass

        exact_matches = []

        for drive in self.get_drives():

            try:

                for root, dirs, files in os.walk(
                    drive,
                    topdown=True,
                ):

                    dirs[:] = [
                        directory
                        for directory in dirs
                        if directory.lower()
                        not in self.SKIP_DIRECTORIES
                    ]

                    for filename in files:

                        if (
                            filename.lower()
                            == file_reference.lower()
                        ):

                            path = (
                                Path(root)
                                / filename
                            )

                            try:

                                if path.is_file():

                                    exact_matches.append(
                                        path
                                    )

                            except OSError:

                                continue

                            if (
                                len(exact_matches)
                                >= 5
                            ):

                                break

                    if (
                        len(exact_matches)
                        >= 5
                    ):

                        break

            except (
                PermissionError,
                OSError,
            ):

                continue

            if (
                len(exact_matches)
                >= 5
            ):

                break

        if len(exact_matches) == 1:

            return self.read_file(
                exact_matches[0]
            )

        if len(exact_matches) > 1:

            self.last_search_results = (
                exact_matches[
                    :self.MAX_RESULTS
                ]
            )

            self.last_search_type = "file"

            return self.format_search_results(
                self.last_search_results,
                file_reference
            )

        partial_matches = (
            self.search_matching_files(
                file_reference
            )
        )

        if len(partial_matches) == 1:

            return self.read_file(
                partial_matches[0]
            )

        if len(partial_matches) > 1:

            self.last_search_results = (
                partial_matches[
                    :self.MAX_RESULTS
                ]
            )

            self.last_search_type = "file"

            return self.format_search_results(
                self.last_search_results,
                file_reference
            )

        return (
            f"Sir, I couldn't find a file named "
            f"'{file_reference}' on the available drives."
        )

    # ==========================================
    # SEARCH MATCHING FILES
    # ==========================================

    def search_matching_files(
        self,
        query
    ):

        query = self.normalize_search_query(
            query
        )

        results = []

        for drive in self.get_drives():

            try:

                for root, dirs, files in os.walk(
                    drive,
                    topdown=True,
                ):

                    dirs[:] = [
                        directory
                        for directory in dirs
                        if directory.lower()
                        not in self.SKIP_DIRECTORIES
                    ]

                    for filename in files:

                        if (
                            query
                            not in filename.lower()
                        ):

                            continue

                        path = (
                            Path(root)
                            / filename
                        )

                        try:

                            if path.is_file():

                                results.append(
                                    path
                                )

                        except OSError:

                            continue

            except (
                PermissionError,
                OSError,
            ):

                continue

        if results:

            query_lower = query.lower()

            results.sort(
                key=lambda p: (
                    0
                    if p.name.lower()
                    == query_lower
                    else 1
                    if p.name.lower()
                    .startswith(query_lower)
                    else 2
                )
            )

        return results[
            :self.MAX_RESULTS
        ]

    # ==========================================
    # FORMAT SEARCH RESULTS
    # ==========================================

    def format_search_results(
        self,
        results,
        query=""
    ):

        console_lines = []

        console_lines.append(
            f"Found {len(results)} matching file(s):"
        )

        console_lines.append("")

        for index, path in enumerate(
            results,
            start=1
        ):

            console_lines.append(
                f"{index}. {path}"
            )

        console_lines.append("")

        console_lines.append(
            "Say 'read file number N' "
            "to read a file."
        )

        console_lines.append(
            "Say 'open file number N' "
            "to open a file."
        )

        console_lines.append(
            "Say 'close file number N' "
            "to close an opened file."
        )

        console_output = "\n".join(
            console_lines
        )

        voice_output = (
            f"Sir, I found {len(results)} "
            "matching files. Please tell me "
            "the file number."
        )

        return (
            f"{console_output}"
            f"|||VOICE|||"
            f"{voice_output}"
        )

    # ==========================================
    # READ FILE
    # ==========================================

    def read_file(
        self,
        file_path
    ):

        path = Path(
            file_path
        )

        self.last_read_file = path

        if not path.exists():

            self.last_read_file = None

            return (
                f"Sir, the file '{path}' "
                "does not exist."
            )

        if not path.is_file():

            self.last_read_file = None

            return (
                f"Sir, '{path}' is not a file."
            )

        try:

            file_size = path.stat().st_size

        except OSError:

            return (
                "Sir, I couldn't inspect "
                "the file size."
            )

        if path.suffix.lower() == ".pdf":

            if file_size > self.MAX_READ_SIZE:

                return (
                    f"Sir, '{path.name}' is too "
                    "large to read safely."
                )

            return self.read_pdf(path)

        if path.suffix.lower() == ".docx":

            if file_size > self.MAX_READ_SIZE:

                return (
                    f"Sir, '{path.name}' is too "
                    "large to read safely."
                )

            return self.read_docx(path)

        if path.suffix.lower() == ".doc":

            return (
                f"Sir, '{path.name}' is an old "
                "DOC format file. DOCX and PDF "
                "reading are supported currently."
            )

        if file_size > self.MAX_READ_SIZE:

            return (
                "Sir, that file is too large "
                "to read safely."
            )

        if (
            path.suffix.lower()
            not in self.TEXT_EXTENSIONS
            and path.name.lower()
            not in {
                ".env",
                ".gitignore",
            }
        ):

            return (
                f"Sir, '{path.name}' does not "
                "appear to be a supported text "
                "or document file."
            )

        try:

            with open(
                path,
                "rb"
            ) as file:

                raw_data = file.read(
                    self.MAX_READ_SIZE + 1
                )

        except PermissionError:

            return (
                "Sir, I don't have permission "
                "to read that file."
            )

        except OSError as e:

            print(
                "JARVIS file read error:",
                e
            )

            return (
                "Sir, I couldn't read that file."
            )

        if len(raw_data) > self.MAX_READ_SIZE:

            return (
                "Sir, that file is too large "
                "to read safely."
            )

        if b"\x00" in raw_data:

            return (
                f"Sir, '{path.name}' appears to "
                "be a binary file, so I won't "
                "display its raw contents."
            )

        try:

            content = raw_data.decode(
                "utf-8"
            )

        except UnicodeDecodeError:

            try:

                content = raw_data.decode(
                    "cp1252"
                )

            except UnicodeDecodeError:

                try:

                    content = raw_data.decode(
                        "latin-1"
                    )

                except UnicodeDecodeError:

                    return (
                        f"Sir, I couldn't decode "
                        f"'{path.name}' as text."
                    )

        return self.format_text_content(
            path,
            file_size,
            content
        )

    # ==========================================
    # READ PDF
    # ==========================================

    def read_pdf(
        self,
        path
    ):

        try:

            from pypdf import PdfReader

        except ImportError:

            return (
                "Sir, PDF reading is not installed. "
                "Please install pypdf."
            )

        try:

            reader = PdfReader(
                str(path),
                strict=False
            )

            pages = []

            for page_number, page in enumerate(
                reader.pages,
                start=1
            ):

                try:

                    text = page.extract_text()

                except Exception as page_error:

                    print(
                        f"JARVIS PDF page "
                        f"{page_number} error:",
                        page_error
                    )

                    text = ""

                if text:

                    pages.append(
                        f"\n--- Page "
                        f"{page_number} ---\n"
                        f"{text}"
                    )

            content = "\n".join(
                pages
            )

            if not content.strip():

                return (
                    f"Sir, I opened '{path.name}', "
                    "but I couldn't extract readable "
                    "text from the PDF."
                )

            try:

                file_size = path.stat().st_size

            except OSError:

                file_size = 0

            return self.format_text_content(
                path,
                file_size,
                content
            )

        except PermissionError:

            return (
                "Sir, I don't have permission "
                "to read that PDF."
            )

        except OSError as e:

            print(
                "JARVIS PDF OS error:",
                e
            )

            return (
                f"Sir, I couldn't access "
                f"'{path.name}'."
            )

        except Exception as e:

            print(
                "JARVIS PDF read error:",
                repr(e)
            )

            return (
                f"Sir, I couldn't read "
                f"'{path.name}'."
            )

    # ==========================================
    # READ DOCX
    # ==========================================

    def read_docx(
        self,
        path
    ):

        try:

            from docx import Document

        except ImportError:

            return (
                "Sir, DOCX reading is not installed. "
                "Please install python-docx."
            )

        try:

            document = Document(
                str(path)
            )

            paragraphs = []

            for paragraph in document.paragraphs:

                text = paragraph.text.strip()

                if text:

                    paragraphs.append(
                        text
                    )

            for table in document.tables:

                for row in table.rows:

                    cells = []

                    for cell in row.cells:

                        text = cell.text.strip()

                        if text:

                            cells.append(
                                text
                            )

                    if cells:

                        paragraphs.append(
                            " | ".join(cells)
                        )

            content = "\n".join(
                paragraphs
            )

            if not content.strip():

                return (
                    f"Sir, '{path.name}' "
                    "does not contain readable text."
                )

            try:

                file_size = path.stat().st_size

            except OSError:

                file_size = 0

            return self.format_text_content(
                path,
                file_size,
                content
            )

        except PermissionError:

            return (
                "Sir, I don't have permission "
                "to read that DOCX file."
            )

        except OSError as e:

            print(
                "JARVIS DOCX OS error:",
                e
            )

            return (
                f"Sir, I couldn't access "
                f"'{path.name}'."
            )

        except Exception as e:

            print(
                "JARVIS DOCX read error:",
                repr(e)
            )

            return (
                f"Sir, I couldn't read "
                f"'{path.name}'."
            )

    # ==========================================
    # FORMAT TEXT CONTENT
    # ==========================================

    def format_text_content(
        self,
        path,
        file_size,
        content
    ):

        truncated = False

        if (
            len(content)
            > self.MAX_DISPLAY_CHARS
        ):

            content = content[
                :self.MAX_DISPLAY_CHARS
            ]

            truncated = True

        output = []

        output.append(
            f"File: {path}"
        )

        output.append(
            f"Size: {file_size} bytes"
        )

        output.append(
            "-" * 42
        )

        output.append(
            content
        )

        if truncated:

            output.append("")

            output.append(
                "[Content truncated for safety.]"
            )

        output.append("")

        output.append(
            "|||VOICE|||"
        )

        voice_output = (
            f"Sir, I have read {path.name}. "
            "The file contents are displayed "
            "in the console."
        )

        output.append(
            voice_output
        )

        return "\n".join(
            output
        )

    # ==========================================
    # INSPECT FILE
    # ==========================================

    def inspect_file(
        self,
        file_reference
    ):

        file_reference = (
            self.normalize_file_reference(
                file_reference
            )
        )

        path = None

        for result in self.last_search_results:

            try:

                if (
                    result.exists()
                    and result.is_file()
                    and result.name.lower()
                    == file_reference.lower()
                ):

                    path = result

                    break

            except OSError:

                continue

        if path is None:

            possible_path = Path(
                file_reference
            )

            try:

                if (
                    possible_path.exists()
                    and possible_path.is_file()
                ):

                    path = possible_path

            except OSError:

                pass

        if (
            path is None
            and self.is_resume_reference(
                file_reference
            )
        ):

            matches = (
                self.search_matching_files(
                    "resume"
                )
            )

            if len(matches) == 1:

                path = matches[0]

        if path is None:

            return (
                f"Sir, I couldn't find "
                f"'{file_reference}' "
                "on the available drives."
            )

        try:

            stats = path.stat()

            return "\n".join([
                f"File: {path}",
                f"Name: {path.name}",
                f"Extension: "
                f"{path.suffix or 'none'}",
                f"Size: {stats.st_size} bytes",
                f"Parent folder: {path.parent}",
            ])

        except OSError:

            return (
                "Sir, I couldn't inspect "
                "that file."
            )

    # ==========================================
    # INSPECT FOLDER
    # ==========================================

    def inspect_folder(
        self,
        folder_reference
    ):

        folder_reference = (
            self.normalize_file_reference(
                folder_reference
            )
        )

        path = None

        for result in self.last_search_results:

            try:

                if (
                    result.exists()
                    and result.is_dir()
                    and result.name.lower()
                    == folder_reference.lower()
                ):

                    path = result

                    break

            except OSError:

                continue

        if path is None:

            possible_path = Path(
                folder_reference
            )

            try:

                if (
                    possible_path.exists()
                    and possible_path.is_dir()
                ):

                    path = possible_path

            except OSError:

                pass

        if path is None:

            matches = (
                self.search_matching_folders(
                    folder_reference
                )
            )

            if len(matches) == 1:

                path = matches[0]

        if path is None:

            return (
                f"Sir, I couldn't find "
                f"the folder '{folder_reference}'."
            )

        try:

            files = 0

            folders = 0

            for item in path.iterdir():

                if item.is_file():

                    files += 1

                elif item.is_dir():

                    folders += 1

            return "\n".join([
                f"Folder: {path}",
                f"Files: {files}",
                f"Folders: {folders}",
            ])

        except PermissionError:

            return (
                "Sir, I don't have permission "
                "to inspect that folder."
            )

        except OSError:

            return (
                "Sir, I couldn't inspect "
                "that folder."
            )

    # ==========================================
    # NORMALIZE SEARCH QUERY
    # ==========================================

    def normalize_search_query(
        self,
        query
    ):

        query = str(
            query
        ).lower().strip()

        query = query.strip(
            "\"'"
        )

        replacements = {

            " dot ": ".",
            " period ": ".",

            " underscore ": "_",

            " hyphen ": "-",
            " dash ": "-",

        }

        for spoken, symbol in replacements.items():

            query = query.replace(
                spoken,
                symbol
            )

        if query in [

            "res",
            "resu",
            "resum",
            "resume",
            "my resume",
            "resume file",

        ]:

            return "resume"

        if query in [

            "pdf file",
            "pdf files",
            "pdf document",
            "pdf documents",

        ]:

            return ".pdf"

        if query in [

            "docx file",
            "docx files",
            "word document",
            "word documents",

        ]:

            return ".docx"

        return query.strip()

    # ==========================================
    # NORMALIZE FILE REFERENCE
    # ==========================================

    def normalize_file_reference(
        self,
        reference
    ):

        reference = str(
            reference
        ).lower().strip()

        reference = reference.strip(
            "\"'"
        )

        replacements = {

            " dot ": ".",
            " period ": ".",

            " underscore ": "_",

            " hyphen ": "-",
            " dash ": "-",

            " space ": " ",

        }

        for spoken, symbol in replacements.items():

            reference = reference.replace(
                spoken,
                symbol
            )

        if reference in [

            "res",
            "resu",
            "resum",
            "resume",
            "my resume",
            "resume file",
            "my resume file",

        ]:

            return "resume"

        extension_replacements = {

            " py": ".py",
            " txt": ".txt",
            " pdf": ".pdf",
            " docx": ".docx",
            " doc": ".doc",
            " csv": ".csv",
            " json": ".json",
            " xml": ".xml",
            " html": ".html",
            " css": ".css",
            " js": ".js",
            " java": ".java",
            " cpp": ".cpp",
            " c": ".c",

        }

        for spoken, extension in (
            extension_replacements.items()
        ):

            if reference.endswith(
                spoken
            ):

                reference = (
                    reference[
                        :-len(spoken)
                    ]
                    + extension
                )

                break

        return reference.strip()

    # ==========================================
    # RESUME REFERENCE CHECK
    # ==========================================

    def is_resume_reference(
        self,
        reference
    ):

        reference = (
            reference.lower()
            .strip()
        )

        return reference in [

            "resume",
            "res",
            "resu",
            "resum",
            "my resume",
            "resume file",
            "my resume file",

        ]

    # ==========================================
    # NUMBER PARSER
    # ==========================================

    def parse_number(
        self,
        value
    ):

        value = str(
            value
        ).lower().strip()

        if value.isdigit():

            return int(value)

        number_words = {

            "zero": 0,

            "one": 1,
            "two": 2,
            "three": 3,
            "four": 4,
            "five": 5,
            "six": 6,
            "seven": 7,
            "eight": 8,
            "nine": 9,

            "ten": 10,
            "eleven": 11,
            "twelve": 12,
            "thirteen": 13,
            "fourteen": 14,
            "fifteen": 15,
            "sixteen": 16,
            "seventeen": 17,
            "eighteen": 18,
            "nineteen": 19,

            "twenty": 20,

        }

        return number_words.get(
            value
        )